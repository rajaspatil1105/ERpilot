"""Assessment report rendering for the educational ERPilot prototype."""

from __future__ import annotations

import json
import textwrap
from collections.abc import Mapping
from datetime import datetime, timezone
from io import BytesIO
from typing import Any

from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas


class ReportDataError(ValueError):
    """Raised when persisted assessment data cannot form a safe report."""


def build_report_data(assessment: Any) -> dict[str, Any]:
    """Validate persisted assessment data and prepare report content."""

    patient = getattr(assessment, "patient", None)
    if patient is None:
        raise ReportDataError("assessment has no associated patient")

    evidence = _load_object(
        getattr(assessment, "evidence_json", None),
        "evidence_json",
    )
    vitals = _load_object(
        getattr(assessment, "vitals_json", None),
        "vitals_json",
    )
    inference = _load_object(
        getattr(assessment, "inference_result_json", None),
        "inference_result_json",
    )
    if inference.get("outcome") not in {
        "FULL_MATCH",
        "PARTIAL_MATCH",
        "NO_MATCH",
    }:
        raise ReportDataError("inference_result_json.outcome is missing or invalid")
    for key in (
        "evaluations",
        "full_matches",
        "partial_matches",
        "no_matches",
        "explanations",
        "follow_up_questions",
    ):
        if not isinstance(inference.get(key), list):
            raise ReportDataError(
                f"inference_result_json.{key} must be a list"
            )
    evaluations = inference["evaluations"]
    explanations = inference["explanations"]
    answered_follow_up_questions = inference.get(
        "answered_follow_up_questions", []
    )
    if not isinstance(answered_follow_up_questions, list):
        raise ReportDataError(
            "inference_result_json.answered_follow_up_questions must be a list"
        )

    modifiers = inference.get("modifiers", {})
    if not isinstance(modifiers, Mapping):
        raise ReportDataError("inference_result_json.modifiers must be an object")
    if modifiers.get("outcome") not in {"FULL_MATCH", "NO_MATCH"}:
        raise ReportDataError(
            "inference_result_json.modifiers.outcome is missing or invalid"
        )
    modifier_evidence = modifiers.get("submitted_evidence")
    if not isinstance(modifier_evidence, Mapping):
        raise ReportDataError(
            "inference_result_json.modifiers.submitted_evidence must be an object"
        )
    for key in ("evaluations", "fired_modifiers", "no_matches"):
        if not isinstance(modifiers.get(key), list):
            raise ReportDataError(
                f"inference_result_json.modifiers.{key} must be a list"
            )
    modifier_evaluations = modifiers["evaluations"]

    return {
        "assessment_id": assessment.id,
        "patient_id": patient.id,
        "category": assessment.category,
        "created_at": _format_datetime(assessment.created_at),
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "patient": {
            "name": patient.name,
            "age": patient.age,
            "gender": patient.gender,
            "incident_type": patient.incident_type,
        },
        "evidence": dict(evidence),
        "vitals": dict(vitals),
        "inference": inference,
        "evaluations": evaluations,
        "explanations": explanations,
        "answered_follow_up_questions": answered_follow_up_questions,
        "modifiers": dict(modifiers),
        "modifier_evidence": dict(modifier_evidence),
        "modifier_evaluations": modifier_evaluations,
    }


def generate_pdf_report(report_data: Mapping[str, Any]) -> bytes:
    """Render a plain, accessible PDF report without changing inference data."""

    output = BytesIO()
    document = canvas.Canvas(output, pagesize=letter, pageCompression=0)
    width, height = letter
    left = 54
    bottom_margin = 54
    y = height - bottom_margin

    def write(text: str, *, heading: bool = False) -> None:
        nonlocal y
        font = "Helvetica-Bold" if heading else "Helvetica"
        size = 13 if heading else 9
        line_height = 17 if heading else 12
        document.setFont(font, size)
        for line in _wrapped_lines(text, 96 if heading else 112):
            if y < bottom_margin:
                document.showPage()
                y = height - bottom_margin
                document.setFont(font, size)
            document.drawString(left, y, _pdf_text(line))
            y -= line_height
        if heading:
            y -= 4

    _write_report_sections(write, report_data)
    document.save()
    return output.getvalue()


def generate_docx_report(report_data: Mapping[str, Any]) -> bytes:
    """Render a DOCX report using the project's supported DOCX library."""

    document = Document()

    def heading(text: str, level: int = 1) -> None:
        document.add_heading(text, level=level)

    def paragraph(text: str = "") -> None:
        document.add_paragraph(_plain_text(text))

    def bullet(text: str) -> None:
        document.add_paragraph(_plain_text(text), style="List Bullet")

    heading("ERPilot AI Assessment Report")
    paragraph(
        "Educational prototype only. Not clinically validated, not intended "
        "for real-patient use, and not a substitute for qualified clinical "
        "judgment."
    )
    paragraph(f"Generated at: {report_data['generated_at']}")

    heading("Assessment and patient", level=2)
    for label, value in _identity_lines(report_data):
        paragraph(f"{label}: {value}")

    heading("Submitted core evidence", level=2)
    paragraph(_json_text(report_data["evidence"]))
    paragraph(f"Vital signs: {_json_text(report_data['vitals'])}")

    inference = report_data["inference"]
    paragraph(f"Core software outcome: {inference.get('outcome', 'Unavailable')}")
    paragraph(
        f"Triggered core rule IDs: "
        f"{_json_text(_matched_rule_ids(inference))}"
    )
    heading("Core rule evaluations", level=2)
    for evaluation in report_data["evaluations"]:
        bullet(_evaluation_text(evaluation))
    for explanation in report_data["explanations"]:
        if isinstance(explanation, Mapping):
            paragraph(
                f"Explanation for {explanation.get('rule_id', 'unknown')}: "
                f"{explanation.get('text', 'Unavailable')}"
            )

    heading("Follow-up questions", level=2)
    if report_data["answered_follow_up_questions"]:
        for answer in report_data["answered_follow_up_questions"]:
            paragraph(_follow_up_text(answer))
    else:
        paragraph("No follow-up answer history is present in the stored record.")
    if inference["follow_up_questions"]:
        paragraph("Currently unanswered questions:")
        for question in inference["follow_up_questions"]:
            paragraph(_follow_up_text(question, unanswered=True))
    else:
        paragraph("No follow-up questions remain unanswered.")

    heading("Special-condition modifiers", level=2)
    paragraph(
        "Modifier results are additive review flags. They do not replace or "
        "rewrite the core trauma outcome, and they are not diagnoses or "
        "treatment recommendations."
    )
    paragraph(
        f"Submitted modifier evidence: "
        f"{_json_text(report_data['modifier_evidence'])}"
    )
    paragraph(
        f"Modifier outcome: "
        f"{report_data['modifiers'].get('outcome', 'NO_MATCH')}"
    )
    paragraph(
        f"Triggered modifier IDs: "
        f"{_json_text(_fired_modifier_ids(report_data['modifiers']))}"
    )
    for evaluation in report_data["modifier_evaluations"]:
        bullet(_modifier_text(evaluation))

    heading("Priority and safety limitation", level=2)
    paragraph(
        "RED, ORANGE, and YELLOW values are provisional ERPilot software "
        "display priorities only. They are not clinically validated triage "
        "categories."
    )
    paragraph(
        "This report reflects recorded evidence and configured rule "
        "evaluations. It does not establish a diagnosis or guarantee a "
        "clinical outcome."
    )

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _write_report_sections(write, report_data: Mapping[str, Any]) -> None:
    write("ERPilot AI Assessment Report", heading=True)
    write(
        "Educational prototype only. Not clinically validated, not intended "
        "for real-patient use, and not a substitute for qualified clinical "
        "judgment."
    )
    write(f"Generated at: {report_data['generated_at']}")

    write("Assessment and patient", heading=True)
    for label, value in _identity_lines(report_data):
        write(f"{label}: {value}")

    write("Submitted core evidence", heading=True)
    write(_json_text(report_data["evidence"]))
    write(f"Vital signs: {_json_text(report_data['vitals'])}")

    inference = report_data["inference"]
    write(f"Core software outcome: {inference.get('outcome', 'Unavailable')}")
    write(f"Triggered core rule IDs: {_json_text(_matched_rule_ids(inference))}")
    write("Core rule evaluations", heading=True)
    for evaluation in report_data["evaluations"]:
        write(_evaluation_text(evaluation))
    for explanation in report_data["explanations"]:
        if isinstance(explanation, Mapping):
            write(
                f"Explanation for {explanation.get('rule_id', 'unknown')}: "
                f"{explanation.get('text', 'Unavailable')}"
            )

    write("Follow-up questions", heading=True)
    if report_data["answered_follow_up_questions"]:
        for answer in report_data["answered_follow_up_questions"]:
            write(_follow_up_text(answer))
    else:
        write("No follow-up answer history is present in the stored record.")
    if inference["follow_up_questions"]:
        write("Currently unanswered questions:")
        for question in inference["follow_up_questions"]:
            write(_follow_up_text(question, unanswered=True))
    else:
        write("No follow-up questions remain unanswered.")

    write("Special-condition modifiers", heading=True)
    write(
        "Modifier results are additive review flags. They do not replace or "
        "rewrite the core trauma outcome, and they are not diagnoses or "
        "treatment recommendations."
    )
    write(f"Submitted modifier evidence: {_json_text(report_data['modifier_evidence'])}")
    write(
        f"Modifier outcome: "
        f"{report_data['modifiers'].get('outcome', 'NO_MATCH')}"
    )
    write(
        f"Triggered modifier IDs: "
        f"{_json_text(_fired_modifier_ids(report_data['modifiers']))}"
    )
    for evaluation in report_data["modifier_evaluations"]:
        write(_modifier_text(evaluation))

    write("Priority and safety limitation", heading=True)
    write(
        "RED, ORANGE, and YELLOW values are provisional ERPilot software "
        "display priorities only. They are not clinically validated triage "
        "categories."
    )
    write(
        "This report reflects recorded evidence and configured rule "
        "evaluations. It does not establish a diagnosis or guarantee a "
        "clinical outcome."
    )


def _identity_lines(report_data: Mapping[str, Any]) -> list[tuple[str, Any]]:
    patient = report_data["patient"]
    return [
        ("Assessment ID", report_data["assessment_id"]),
        ("Patient ID", report_data["patient_id"]),
        ("Patient name", patient["name"]),
        ("Age", patient["age"]),
        ("Gender", patient["gender"]),
        ("Incident type", patient["incident_type"]),
        ("Trauma category", report_data["category"]),
        ("Assessment created at", report_data["created_at"]),
    ]


def _evaluation_text(evaluation: Any) -> str:
    if not isinstance(evaluation, Mapping):
        return "Malformed core rule evaluation."
    return (
        f"Core rule {evaluation.get('rule_id', 'unknown')}: "
        f"{evaluation.get('match_status', 'UNKNOWN')}; "
        f"priority={evaluation.get('priority_flag', 'Unavailable')}; "
        f"conditions={_json_text(evaluation.get('present_conditions', []))}; "
        f"checklist={evaluation.get('checklist_action', 'Unavailable')}; "
        f"protocol={evaluation.get('protocol_reference', 'Unavailable')}"
    )


def _matched_rule_ids(inference: Mapping[str, Any]) -> list[Any]:
    return [
        evaluation.get("rule_id")
        for key in ("full_matches", "partial_matches")
        for evaluation in inference.get(key, [])
        if isinstance(evaluation, Mapping)
    ]


def _fired_modifier_ids(modifiers: Mapping[str, Any]) -> list[Any]:
    return [
        evaluation.get("modifier_id")
        for evaluation in modifiers.get("fired_modifiers", [])
        if isinstance(evaluation, Mapping)
    ]


def _modifier_text(evaluation: Any) -> str:
    if not isinstance(evaluation, Mapping):
        return "Malformed modifier evaluation."
    return (
        f"Modifier {evaluation.get('modifier_id', 'unknown')}: "
        f"{evaluation.get('match_status', 'UNKNOWN')}; "
        f"action={evaluation.get('action', 'Unavailable')}; "
        f"priority override={evaluation.get('priority_override', 'Unavailable')}; "
        f"protocol={evaluation.get('protocol_reference', 'Unavailable')}; "
        f"explanation={evaluation.get('explanation', 'Unavailable')}"
    )


def _follow_up_text(question: Any, *, unanswered: bool = False) -> str:
    if not isinstance(question, Mapping):
        return "Malformed follow-up question record."
    question_id = question.get("question_id", "unknown")
    prompt = question.get("prompt", "Unavailable")
    evidence_key = question.get("evidence_key", "Unavailable")
    if unanswered:
        return (
            f"{question_id}: {prompt}; evidence={evidence_key}; "
            "response=UNANSWERED"
        )
    return (
        f"{question_id}: {prompt}; evidence={evidence_key}; "
        f"response={question.get('response', 'Unavailable')}; "
        f"recorded state={question.get('recorded_state', 'Unavailable')}"
    )


def _load_object(value: Any, field_name: str) -> dict[str, Any]:
    if not isinstance(value, str):
        raise ReportDataError(f"{field_name} must contain serialized JSON")
    try:
        parsed = json.loads(value)
    except json.JSONDecodeError as error:
        raise ReportDataError(f"{field_name} contains malformed JSON") from error
    if not isinstance(parsed, Mapping):
        raise ReportDataError(f"{field_name} must contain a JSON object")
    return dict(parsed)


def _format_datetime(value: Any) -> str:
    if isinstance(value, datetime):
        return value.isoformat()
    if value is None:
        return "Unavailable"
    return _plain_text(value)


def _json_text(value: Any) -> str:
    try:
        return json.dumps(value, sort_keys=True, ensure_ascii=True)
    except (TypeError, ValueError):
        return _plain_text(value)


def _wrapped_lines(value: Any, width: int) -> list[str]:
    text = _plain_text(value)
    return textwrap.wrap(text, width=width) or [""]


def _plain_text(value: Any) -> str:
    return str(value).replace("\r", " ").replace("\n", " ")


def _pdf_text(value: Any) -> str:
    return _plain_text(value).encode("latin-1", "replace").decode("latin-1")
